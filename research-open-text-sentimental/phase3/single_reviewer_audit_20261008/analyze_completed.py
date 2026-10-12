"""Analyze one blind re-review against the earlier recorded Phase 2 decisions."""

import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "submission/Manu_Blind_Consistency_Review_120_Completed.xlsx"
BLIND = json.loads((HERE / "blind_review_rows.json").read_text())
MAPPING = json.loads((HERE / "ANALYST_ONLY_MAPPING.json").read_text())
BY_ID = {r["audit_id"]: r for r in MAPPING}
EXPECTED_TEXT = {r["audit_id"]: r for r in BLIND}
FIELDS = ["about_software", "update_related", "discussion_type"]
ALLOWED = {
    "about_software": {"yes", "no", "unclear"},
    "update_related": {"yes", "no", "unclear"},
    "discussion_type": {"technical_troubleshooting", "general_dissatisfaction", "other", "na"},
    "confidence": {"high", "medium", "low"},
}

wb = load_workbook(SOURCE, read_only=True, data_only=True)
assert wb.sheetnames == ["Review"]
sheet = wb["Review"]
header = [c.value for c in sheet[5][:9]]
assert header == ["audit_id", "title", "post_body", "saved_reply_context", "about_software", "update_related", "discussion_type", "confidence", "reason"]
rows = list(sheet.iter_rows(min_row=6, max_row=125, min_col=1, max_col=9, values_only=True))
assert len(rows) == len(BY_ID) == 120

submitted = []
for values in rows:
    audit_id = values[0]
    assert audit_id in BY_ID, f"Unknown ID: {audit_id}"
    expected = EXPECTED_TEXT[audit_id]
    for col, val in zip(("title", "post_body", "saved_reply_context"), values[1:4]):
        assert (val or "") == (expected[col] or ""), f"Post text changed: {audit_id} {col}"
    answer = dict(zip(header, values))
    for field, valid in ALLOWED.items():
        assert answer[field] in valid, f"Invalid or blank {field}: {audit_id}={answer[field]!r}"
    assert isinstance(answer["reason"], str) and answer["reason"].strip(), f"Missing reason: {audit_id}"
    assert (answer["discussion_type"] == "na") == (answer["about_software"] != "yes"), f"Discussion type incompatible with scope: {audit_id}"
    submitted.append(answer)
assert len({r["audit_id"] for r in submitted}) == 120

matrix = {field: Counter() for field in FIELDS}
by_stratum = defaultdict(lambda: {field: Counter() for field in FIELDS})
disagreements = []
full_match = 0
for answer in submitted:
    previous = BY_ID[answer["audit_id"]]
    group = previous["sampling_stratum_audit_only"]
    changed = []
    for field in FIELDS:
        old = previous[f"previous_{field}_audit_only"]
        new = answer[field]
        matrix[field][(old, new)] += 1
        by_stratum[group][field][(old, new)] += 1
        if old != new:
            changed.append(field)
    if not changed:
        full_match += 1
    else:
        disagreements.append({
            "audit_id": answer["audit_id"],
            "reddit_id": previous["reddit_id"],
            "sampling_stratum": group,
            "changed_fields": ",".join(changed),
            "previous_about_software": previous["previous_about_software_audit_only"],
            "new_about_software": answer["about_software"],
            "previous_update_related": previous["previous_update_related_audit_only"],
            "new_update_related": answer["update_related"],
            "previous_discussion_type": previous["previous_discussion_type_audit_only"],
            "new_discussion_type": answer["discussion_type"],
            "new_confidence": answer["confidence"],
            "new_reason": answer["reason"],
            "title": answer["title"],
        })

def summarize(counter):
    return {
        "n": sum(counter.values()),
        "matches": sum(n for (old, new), n in counter.items() if old == new),
        "confusion": [{"previous": old, "new": new, "n": n} for (old, new), n in sorted(counter.items())],
    }

results = {
    "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "sample_n": len(submitted),
    "all_three_fields_match_n": full_match,
    "any_field_disagreement_n": len(disagreements),
    "new_confidence_counts": dict(Counter(r["confidence"] for r in submitted)),
    "fields": {field: summarize(matrix[field]) for field in FIELDS},
    "by_prior_stratum": {group: {field: summarize(counters[field]) for field in FIELDS} for group, counters in by_stratum.items()},
    "provenance_note": "One new blind review versus prior mixed-provenance recorded decisions. Not independent inter-annotator agreement or a natural-prevalence estimate.",
}
(HERE / "submission/analysis.json").write_text(json.dumps(results, indent=2) + "\n")
with (HERE / "submission/disagreements.csv").open("w", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=list(disagreements[0]) if disagreements else ["audit_id"])
    writer.writeheader()
    writer.writerows(disagreements)
print(json.dumps({"sample_n": len(submitted), "all_match": full_match, "disagreements": len(disagreements), "fields": {k: {"n": v["n"], "matches": v["matches"]} for k, v in results["fields"].items()}}, indent=2))
