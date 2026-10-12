"""Validate and preserve the completed 22-row final-decision workbook."""

import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "submission/Adjudicate_22_Differences_Completed.xlsx"
EXPECTED = json.loads((HERE / "disagreement_rows.json").read_text())
BY_ID = {r["audit_id"]: r for r in EXPECTED}
HEADERS = ["audit_id", "title", "post_body", "saved_reply_context", "previous_about_software", "new_about_software", "final_about_software", "previous_update_related", "new_update_related", "final_update_related", "previous_discussion_type", "new_discussion_type", "final_discussion_type", "final_reason"]
ALLOWED = {"final_about_software": {"yes", "no", "unclear"}, "final_update_related": {"yes", "no", "unclear"}, "final_discussion_type": {"technical_troubleshooting", "general_dissatisfaction", "other", "na"}}

wb = load_workbook(SOURCE, read_only=True, data_only=True)
assert wb.sheetnames == ["Review"]
sheet = wb["Review"]
assert [c.value for c in sheet[5][:14]] == HEADERS
observed = [dict(zip(HEADERS, values)) for values in sheet.iter_rows(min_row=6, max_row=27, min_col=1, max_col=14, values_only=True)]
assert len(observed) == len(BY_ID) == len({r["audit_id"] for r in observed}) == 22
final = []
for row in observed:
    ident = row["audit_id"]
    assert ident in BY_ID
    original = BY_ID[ident]
    for key in HEADERS:
        if key.startswith("final_"):
            continue
        assert (row[key] or "") == (original[key] or ""), f"Previously supplied content changed: {ident} {key}"
    for key, allowed in ALLOWED.items():
        assert row[key] in allowed, f"Missing/invalid {key}: {ident}"
    assert isinstance(row["final_reason"], str) and row["final_reason"].strip(), ident
    assert (row["final_about_software"] != "yes") == (row["final_discussion_type"] == "na"), ident
    final.append({k: row[k] for k in ("audit_id", "final_about_software", "final_update_related", "final_discussion_type", "final_reason")})

path = HERE / "submission/final_decisions.json"
path.write_text(json.dumps(final, ensure_ascii=False, indent=2) + "\n")
summary = {
    "workbook_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "decision_count": len(final),
    "all_fields_final_equal_previous_n": sum(r["final_about_software"] == r["previous_about_software"] and r["final_update_related"] == r["previous_update_related"] and r["final_discussion_type"] == r["previous_discussion_type"] for r in observed),
    "all_fields_final_equal_new_n": sum(r["final_about_software"] == r["new_about_software"] and r["final_update_related"] == r["new_update_related"] and r["final_discussion_type"] == r["new_discussion_type"] for r in observed),
    "update_transitions_from_previous": dict(Counter(f"{r['previous_update_related']}->{r['final_update_related']}" for r in observed)),
    "software_transitions_from_previous": dict(Counter(f"{r['previous_about_software']}->{r['final_about_software']}" for r in observed)),
    "discussion_transitions_from_previous": dict(Counter(f"{r['previous_discussion_type']}->{r['final_discussion_type']}" for r in observed)),
    "master_updated": False,
    "model_retrained": False,
}
(HERE / "submission/ADJUDICATION_SUMMARY.json").write_text(json.dumps(summary, indent=2) + "\n")
print(json.dumps(summary, indent=2))
