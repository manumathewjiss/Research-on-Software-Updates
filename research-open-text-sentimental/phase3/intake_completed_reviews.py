"""Archive and validate completed diagnostic and fresh-test review workbooks."""
from __future__ import annotations

import hashlib
import json
import shutil
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "completed_reviews_20261007_corrected"
INPUTS = {
    "manu_test": Path("/Users/manumathewjiss/Documents/Codex/2026-10-06/i-wi/outputs/Manu_Fresh_Test_300_Completed.xlsx"),
    "berhe_test": Path("/Users/manumathewjiss/Documents/Codex/2026-10-06/i-wi/outputs/Berhe_Fresh_Test_300_Completed.xlsx"),
    "diagnostic": Path("/Users/manumathewjiss/Documents/Codex/2026-10-06/i-wi/outputs/Main_Diagnostic_All_250_Completed.xlsx"),
}
SOURCES = {
    "manu_test": ROOT / "fresh_test_20261007/reviewer_A_rows.json",
    "berhe_test": ROOT / "fresh_test_20261007/reviewer_B_rows.json",
    "diagnostic": ROOT / "diagnostic_pilot/main_sample_20261007",
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_workbook(path: Path) -> tuple[list[str], list[dict]]:
    sheet = openpyxl.load_workbook(path, data_only=True)["Review"]
    heads = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    data = [dict(zip(heads, [sheet.cell(row, col).value for col in range(1, sheet.max_column + 1)]))
            for row in range(2, sheet.max_row + 1)]
    return heads, data


def source_rows(name: str) -> list[dict]:
    path = SOURCES[name]
    if name == "diagnostic":
        return [r for i in range(1, 6) for r in json.loads((path / f"review_batch_{i:02d}.json").read_text())]
    return json.loads(path.read_text())


def validate(name: str, heads: list[str], data: list[dict]) -> dict:
    source = source_rows(name)
    id_field = "sample_id" if name == "diagnostic" else "test_id"
    expected = {r[id_field]: r for r in source}
    actual = {r[id_field]: r for r in data}
    issues = []
    if len(data) != len(source) or len(actual) != len(data) or set(actual) != set(expected):
        issues.append("ID set, row count, or uniqueness mismatch")
    text_changes = [r[id_field] for r in data if r[id_field] in expected and
                    (r["title"] != expected[r[id_field]]["title"] or
                     (r["post_body"] or "") != expected[r[id_field]]["post_body"])]
    if text_changes:
        issues.append(f"Post text changed on {len(text_changes)} rows")
    if name == "diagnostic":
        assert heads == ["sample_id", "title", "post_body", "troubleshooting_check", "software_product",
                         "identifiable_release", "environment", "specific_symptom", "reproduction_context",
                         "error_evidence", "attempted_remedy", "help_request", "evidence_or_note"]
        fields = heads[4:12]
        for r in data:
            check = r["troubleshooting_check"]
            if check not in {"yes", "no", "unclear"}:
                issues.append(f"{r[id_field]} invalid troubleshooting_check")
            if check == "yes" and any(r[f] not in {"present", "not_observed", "unclear"} for f in fields):
                issues.append(f"{r[id_field]} missing/invalid detail field")
            if check != "yes" and any(r[f] is not None for f in fields):
                issues.append(f"{r[id_field]} details filled despite no/unclear check")
            if check != "yes" and not str(r["evidence_or_note"] or "").strip():
                issues.append(f"{r[id_field]} no/unclear check without explanation")
        return {"expected_rows": len(source), "received_rows": len(data), "text_changed_ids": text_changes,
                "issues": issues, "troubleshooting_check": dict(Counter(r["troubleshooting_check"] for r in data)),
                "notes_present_n": sum(bool(str(r["evidence_or_note"] or "").strip()) for r in data),
                "field_counts_among_yes": {f: dict(Counter(r[f] for r in data if r["troubleshooting_check"] == "yes")) for f in fields}}
    assert heads == ["test_id", "title", "post_body", "about_software", "update_related", "discussion_type", "confidence", "reason"]
    for r in data:
        if r["about_software"] not in {"yes", "no", "unclear"} or r["update_related"] not in {"yes", "no", "unclear"}:
            issues.append(f"{r[id_field]} invalid scope field")
        if r["discussion_type"] not in {"technical_troubleshooting", "general_dissatisfaction", "other", "na"}:
            issues.append(f"{r[id_field]} invalid discussion type")
        if r["about_software"] == "yes" and r["discussion_type"] == "na":
            issues.append(f"{r[id_field]} software=yes but type=na")
        if r["about_software"] != "yes" and r["discussion_type"] != "na":
            issues.append(f"{r[id_field]} software not yes but type not na")
        if r["confidence"] not in {"high", "medium", "low"}:
            issues.append(f"{r[id_field]} invalid confidence")
        if not str(r["reason"] or "").strip():
            issues.append(f"{r[id_field]} missing reason")
    return {"expected_rows": len(source), "received_rows": len(data), "text_changed_ids": text_changes,
            "issues": issues, "about_software": dict(Counter(r["about_software"] for r in data)),
            "update_related": dict(Counter(r["update_related"] for r in data)),
            "discussion_type": dict(Counter(r["discussion_type"] for r in data))}


def main() -> None:
    OUT.mkdir(exist_ok=True)
    data_by_name = {}
    summary = {}
    for name, src in INPUTS.items():
        dest = OUT / src.name
        if dest.exists() and sha(dest) != sha(src):
            raise RuntimeError(f"Existing archived workbook differs: {dest}")
        shutil.copy2(src, dest)
        heads, data = read_workbook(src)
        data_by_name[name] = data
        report = validate(name, heads, data)
        report["sha256"] = sha(src)
        summary[name] = report
        (OUT / f"{name}_answers.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")

    a = {r["test_id"]: r for r in data_by_name["manu_test"]}
    b = {r["test_id"]: r for r in data_by_name["berhe_test"]}
    shared = sorted(set(a) & set(b))
    fields = ["about_software", "update_related", "discussion_type", "confidence", "reason"]
    disagreements = []
    for test_id in shared:
        differences = {f: [a[test_id][f], b[test_id][f]] for f in fields if a[test_id][f] != b[test_id][f]}
        if differences:
            disagreements.append({"test_id": test_id, "differences": differences})
    summary["shared_test"] = {
        "n": len(shared),
        "same_on_all_three_label_fields": sum(all(a[i][f] == b[i][f] for f in fields[:3]) for i in shared),
        "same_on_each_field": {f: sum(a[i][f] == b[i][f] for i in shared) for f in fields},
        "same_reason_count": sum(a[i]["reason"] == b[i]["reason"] for i in shared),
        "label_disagreement_ids": [i for i in shared if any(a[i][f] != b[i][f] for f in fields[:3])],
        "any_difference_rows": len(disagreements),
    }
    (OUT / "intake_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n")
    (OUT / "shared_differences.json").write_text(json.dumps(disagreements, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({name: {k:v for k,v in report.items() if k in {"received_rows","issues","troubleshooting_check","about_software","update_related","discussion_type"}}
                      for name, report in summary.items() if name != "shared_test"}, indent=2))
    print(json.dumps(summary["shared_test"], indent=2))


if __name__ == "__main__":
    main()
