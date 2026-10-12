"""Validate and archive a completed single-reviewer pilot without changing labels."""

from __future__ import annotations

import hashlib
import json
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
SUBMITTED = Path("/Users/manumathewjiss/Documents/Codex/2026-10-06/i-wi/outputs/Pilot_Diagnostic_Details_50_Completed.xlsx")
OUT = HERE / "submissions/2026-10-07-manu"
FIELDS = ["software_product", "release_or_update", "environment", "specific_symptom", "trigger_or_steps", "error_evidence", "attempted_fix", "result_after_attempt", "help_request"]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    expected = {row["pilot_id"]: row for row in json.loads((HERE / "pilot_review_text.json").read_text())}
    wb = load_workbook(SUBMITTED, read_only=True, data_only=True)
    assert set(wb.sheetnames) == {"Start here", "Pilot"}
    values = wb["Pilot"].iter_rows(values_only=True)
    headers = list(next(values))
    assert headers == ["pilot_id", "title", "post_body", "troubleshooting_check", *FIELDS, "evidence_or_note"]
    rows = []
    for values_row in values:
        if not values_row[0]:
            continue
        row = dict(zip(headers, values_row))
        pilot_id = row["pilot_id"]
        assert pilot_id in expected
        assert row["title"] == expected[pilot_id]["title"]
        assert (row["post_body"] or "") == expected[pilot_id]["post_body"]
        answer = row["troubleshooting_check"]
        assert answer in {"yes", "no", "unclear"}
        assert isinstance(row["evidence_or_note"], str) and row["evidence_or_note"].strip()
        if answer == "yes":
            for field in FIELDS:
                allowed = {"present", "not_observed", "unclear"}
                if field == "result_after_attempt":
                    allowed.add("not_applicable")
                assert row[field] in allowed, (pilot_id, field, row[field])
        else:
            assert all(row[field] in (None, "") for field in FIELDS), pilot_id
        rows.append(row)
    assert len(rows) == len(expected) == 50
    assert len({row["pilot_id"] for row in rows}) == 50
    assert {row["pilot_id"] for row in rows} == set(expected)

    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SUBMITTED, OUT / SUBMITTED.name)
    answers = [{"pilot_id": row["pilot_id"], "troubleshooting_check": row["troubleshooting_check"], **{field: row[field] for field in FIELDS}, "evidence_or_note": row["evidence_or_note"]} for row in rows]
    (OUT / "answers.json").write_text(json.dumps(answers, indent=2, ensure_ascii=False) + "\n")
    included = [row for row in rows if row["troubleshooting_check"] == "yes"]
    field_counts = {field: dict(Counter(row[field] for row in included)) for field in FIELDS}
    summary = {
        "submitted_sha256": sha(OUT / SUBMITTED.name),
        "pilot_n": len(rows),
        "troubleshooting_check": dict(Counter(row["troubleshooting_check"] for row in rows)),
        "review_flags": [row["pilot_id"] for row in rows if row["troubleshooting_check"] != "yes"],
        "field_counts_among_yes": field_counts,
        "attempt_result_pairs_among_yes": {f"{k[0]}|{k[1]}": v for k, v in Counter((row["attempted_fix"], row["result_after_attempt"]) for row in included).items()},
        "purpose": "Single-reviewer feasibility pilot; quotas are not population-representative",
        "independent_annotator_agreement": None,
        "source_labels_changed": False,
    }
    (OUT / "intake_summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(json.dumps({"pilot_n": len(rows), "troubleshooting_check": summary["troubleshooting_check"], "review_flags": summary["review_flags"], "attempt_result_pairs": summary["attempt_result_pairs_among_yes"]}, indent=2))


if __name__ == "__main__":
    main()
