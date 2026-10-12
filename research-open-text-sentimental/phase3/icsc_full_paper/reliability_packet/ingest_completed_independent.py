"""Validate and archive the two reported independent first-pass submissions.

The user confirms that the reviewers did not discuss or see one another's
answers before saving. The script validates the submitted files, but cannot
independently establish their creation process. It changes no historical data.
"""

from __future__ import annotations

import hashlib
import json
import shutil
from collections import Counter
from pathlib import Path

import openpyxl


HERE = Path(__file__).resolve().parent
INPUTS = Path("/Users/manumathewjiss/Documents/Codex/2026-10-06/i-wi/outputs")
SUBMISSION = HERE / "submissions" / "2026-10-07-independent"
BLIND = json.loads((HERE / "blind_rows.json").read_text())
PRIVATE = json.loads((HERE.parent / "ANALYST_ONLY_RELIABILITY_MAPPING.json").read_text())
FIELDS = [
    "about_software", "update_related", "discussion_type", "confidence", "reason",
    "software_product", "identifiable_release", "environment", "specific_symptom",
    "reproduction_context", "error_evidence", "attempted_remedy", "help_request",
]
ALLOWED = {
    "about_software": {"yes", "no", "unclear"},
    "update_related": {"yes", "no", "unclear"},
    "discussion_type": {"technical_troubleshooting", "general_dissatisfaction", "other", "na"},
    "confidence": {"high", "medium", "low"},
}
DETAILS = set(FIELDS[5:])
for field in DETAILS:
    ALLOWED[field] = {"present", "not_observed", "unclear"}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load(letter: str) -> tuple[Path, list[dict]]:
    source = INPUTS / f"Reviewer_{letter}_200_Completed_Draft.xlsx"
    wb = openpyxl.load_workbook(source, read_only=False, data_only=False)
    assert wb.sheetnames == ["Review"]
    sheet = wb["Review"]
    assert sheet.max_row == 205 and sheet.max_column == 17
    assert [sheet.cell(5, col).value for col in range(1, 18)] == ["review_id", "title", "post_body", "discussion_context", *FIELDS]
    out = []
    for idx, expected in enumerate(BLIND, start=6):
        vals = [sheet.cell(idx, col).value for col in range(1, 18)]
        observed = ["" if value is None else value for value in vals[:4]]
        assert observed == [expected[k] for k in ("review_id", "title", "post_body", "discussion_context")], (letter, idx, "input text changed")
        answer = dict(zip(FIELDS, vals[4:]))
        for field in FIELDS[:4]:
            assert answer[field] in ALLOWED[field], (letter, idx, field, answer[field])
        assert isinstance(answer["reason"], str) and answer["reason"].strip(), (letter, idx, "reason missing")
        if answer["about_software"] != "yes":
            assert answer["discussion_type"] == "na", (letter, idx, "type/scope conflict")
        else:
            assert answer["discussion_type"] != "na", (letter, idx, "type/scope conflict")
        for field in DETAILS:
            if answer["discussion_type"] == "technical_troubleshooting":
                assert answer[field] in ALLOWED[field], (letter, idx, field, answer[field])
            else:
                assert answer[field] in (None, ""), (letter, idx, field, "must be blank")
        out.append({"review_id": expected["review_id"], **answer})
    assert len(out) == len({x["review_id"] for x in out}) == 200
    return source, out


def main() -> None:
    source_a, a = load("A")
    source_b, b = load("B")
    differing_main = Counter()
    differing_reason = []
    for x, y in zip(a, b):
        assert x["review_id"] == y["review_id"]
        for field in FIELDS:
            if x[field] != y[field]:
                if field == "reason":
                    differing_reason.append(x["review_id"])
                else:
                    differing_main[field] += 1
    assert not differing_main, differing_main
    mapping = {r["review_id"]: r for r in PRIVATE}
    assert len(mapping) == 200
    prior_differences = Counter()
    for row in a:
        prior = mapping[row["review_id"]]
        for field in ("about_software", "update_related", "discussion_type"):
            if row[field] != prior[f"prior_{field}"]:
                prior_differences[field] += 1
    summary = {
        "provenance": "Manu reports that both reviewers completed separate first passes without discussing posts or seeing each other's answers before both files were saved. The file contents alone cannot verify the process; earlier corpus exposure remains possible.",
        "status": "validated_reported_independent_first_passes_not_merged",
        "n": 200,
        "source_sha256": {"A": sha(source_a), "B": sha(source_b)},
        "categorical_differences_between_files": dict(differing_main),
        "rows_with_different_reason_text": len(differing_reason),
        "about_software": dict(Counter(r["about_software"] for r in a)),
        "update_related": dict(Counter(r["update_related"] for r in a)),
        "discussion_type": dict(Counter(r["discussion_type"] for r in a)),
        "confidence": dict(Counter(r["confidence"] for r in a)),
        "prior_phase2_label_differences_on_enriched_sample": dict(prior_differences),
        "detail_present_counts_among_agreed_troubleshooting": {field: sum(r[field] == "present" for r in a if r["discussion_type"] == "technical_troubleshooting") for field in FIELDS[5:]},
        "detail_unclear_counts_among_agreed_troubleshooting": {field: sum(r[field] == "unclear" for r in a if r["discussion_type"] == "technical_troubleshooting") for field in FIELDS[5:]},
        "master_changed": False,
        "test_labels_changed": False,
        "independent_agreement_available_under_user_report": True,
    }
    SUBMISSION.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_a, SUBMISSION / source_a.name)
    shutil.copy2(source_b, SUBMISSION / source_b.name)
    (SUBMISSION / "reviewer_a_decisions.json").write_text(json.dumps(a, ensure_ascii=False, indent=2) + "\n")
    (SUBMISSION / "intake_summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
