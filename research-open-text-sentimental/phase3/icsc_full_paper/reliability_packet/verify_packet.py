"""Read-only workbook checks, then record verified hashes in the packet manifest."""

import hashlib
import json
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROWS = json.loads((HERE / "blind_rows.json").read_text())
MANIFEST = HERE / "MANIFEST.json"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def data(letter: str) -> list[list]:
    path = HERE / f"Reviewer_{letter}_Independent_200_Blank.xlsx"
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    ws = wb["Review"]
    assert ws.max_row == 205 and ws.max_column == 17
    assert len(ws.data_validations.dataValidation) == 4
    values = [[ws.cell(r, c).value for c in range(1, 18)] for r in range(6, 206)]
    assert all(row[4:] == [None] * 13 for row in values)
    assert len({row[0] for row in values}) == 200
    assert [tuple("" if value is None else value for value in row[:4]) for row in values] == [
        (x["review_id"], x["title"], x["post_body"], x["discussion_context"]) for x in ROWS
    ]
    return values


assert data("A") == data("B")
manifest = json.loads(MANIFEST.read_text())
assert manifest["selected_n"] == 200
assert manifest["reserved_overlap"] == manifest["fresh_overlap"] == manifest["diagnostic_overlap"] == 0
manifest["reviewer_guide_sha256"] = digest(HERE / "REVIEWER_GUIDE.md")
manifest["reviewer_a_workbook_sha256"] = digest(HERE / "Reviewer_A_Independent_200_Blank.xlsx")
manifest["reviewer_b_workbook_sha256"] = digest(HERE / "Reviewer_B_Independent_200_Blank.xlsx")
manifest["verified_answer_cells_blank"] = True
manifest["verified_reviewer_inputs_identical"] = True
MANIFEST.write_text(json.dumps(manifest, indent=2) + "\n")
print("Verified 200 unique identical input rows and 13 blank answer columns in both workbooks; no test or diagnostic overlap.")
